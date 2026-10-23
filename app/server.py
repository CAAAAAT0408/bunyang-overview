"""분양개요 자동작성 — 사내망용 작은 웹서버.

  python server.py            # 사내망 공개 (접속키 필요)
  python server.py --local    # 이 PC 에서만

이 PC 에서는 http://localhost:8791/ 로, 다른 PC 에서는 http://<이 PC IP>:8791/?k=<접속키> 로 들어온다.
접속키는 app\\_key.txt 에 있다. 엑셀 작성은 이 PC 의 Excel 로 하므로 한 번에 하나씩 처리한다.
"""
import json
import logging
import os
import random
import re
import shutil
import socket
import sys
import threading
import time
import traceback
import urllib.parse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

APP = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, APP)

PORT = 8791
JOBS = os.path.join(APP, '_jobs')
KEYFILE = os.path.join(APP, '_key.txt')
URLFILE = os.path.join(APP, '_url.txt')
LOGFILE = os.path.join(APP, '_server.log')
MAX_UPLOAD = 80 * 1024 * 1024
KEEP_DAYS = 3
EXCEL_LOCK = threading.Lock()

os.makedirs(JOBS, exist_ok=True)
logging.basicConfig(filename=LOGFILE, level=logging.INFO, encoding='utf-8',
                    format='%(asctime)s %(levelname)s %(message)s')
log = logging.getLogger('server')


class UserError(Exception):
    pass


def get_key():
    if os.path.exists(KEYFILE):
        k = open(KEYFILE, encoding='utf-8').read().strip()
        if k:
            return k
    k = ''.join(random.choice('abcdefghijkmnpqrstuvwxyz23456789') for _ in range(8))
    open(KEYFILE, 'w', encoding='utf-8').write(k)
    return k


KEY = get_key()


def lan_ip():
    ips = []
    try:
        ips = socket.gethostbyname_ex(socket.gethostname())[2]
    except Exception:
        pass
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('10.255.255.255', 1))
        ips.insert(0, s.getsockname()[0])
        s.close()
    except Exception:
        pass
    for ip in ips:
        if re.match(r'^(10\.|192\.168\.|172\.(1[6-9]|2\d|3[01])\.)', ip):
            return ip
    return ips[0] if ips else '127.0.0.1'


def cleanup_jobs():
    now = time.time()
    for d in os.listdir(JOBS):
        p = os.path.join(JOBS, d)
        try:
            if os.path.isdir(p) and now - os.path.getmtime(p) > KEEP_DAYS * 86400:
                shutil.rmtree(p, ignore_errors=True)
        except Exception:
            pass


def job_dir(job):
    if not re.fullmatch(r'[0-9a-f]{12}', job or ''):
        raise UserError('작업 번호가 올바르지 않습니다.')
    d = os.path.join(JOBS, job)
    if not os.path.isdir(d):
        raise UserError('작업이 만료되었습니다. 파일을 다시 올려 주세요.')
    return d


def rjson(path, default=None):
    try:
        with open(path, encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return default


def wjson(path, obj):
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, default=str)
    # 화면이 진행 기록(state.json)을 읽는 순간과 겹치면 Windows 가 바꾸기를 막는다(WinError 5) → 잠깐 뒤 다시
    for i in range(40):
        try:
            os.replace(tmp, path)
            return
        except PermissionError:
            if i == 39:
                raise
            time.sleep(0.05)


def _dec(b):
    for enc in ('utf-8', 'cp949'):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            pass
    return b.decode('utf-8', 'replace')


def parse_multipart(ctype, raw):
    m = re.search(r'boundary="?([^";]+)"?', ctype)
    if not m:
        raise UserError('업로드 형식이 올바르지 않습니다.')
    sep = b'--' + m.group(1).encode('latin-1')
    files, form = {}, {}
    for part in raw.split(sep):
        if not part or part.startswith(b'--'):
            continue
        part = part[2:] if part.startswith(b'\r\n') else part
        head, _, data = part.partition(b'\r\n\r\n')
        if data.endswith(b'\r\n'):
            data = data[:-2]
        hd = _dec(head)
        nm = re.search(r'\bname="([^"]*)"', hd)
        fn = re.search(r"filename\*=UTF-8''([^;\r\n]+)", hd, re.I)
        fname = urllib.parse.unquote(fn.group(1)) if fn else None
        if fname is None:
            fn = re.search(r'filename="([^"]*)"', hd)
            fname = fn.group(1) if fn else None
        if not nm:
            continue
        if fname is not None:
            files[nm.group(1)] = (fname.replace('\\', '/').split('/')[-1], data)
        else:
            form[nm.group(1)] = _dec(data)
    return files, form


# ------------------------------------------------------------------ 작업 (분석 · 수집 · 작성)

class Job:
    def __init__(self, job):
        self.id = job
        self.dir = job_dir(job)
        self.state_path = os.path.join(self.dir, 'state.json')

    def state(self):
        return rjson(self.state_path, {'status': 'unknown', 'log': []})

    def set(self, **kw):
        st = self.state()
        st.update(kw)
        wjson(self.state_path, st)

    def log(self, msg):
        st = self.state()
        st.setdefault('log', []).append(f"{time.strftime('%H:%M:%S')} {msg}")
        st['log'] = st['log'][-200:]
        wjson(self.state_path, st)
        log.info('%s %s', self.id, msg)

    @property
    def img_dir(self):
        d = os.path.join(self.dir, 'img')
        os.makedirs(d, exist_ok=True)
        return d

    def cands(self):
        return rjson(os.path.join(self.dir, 'cands.json'), [])

    def save_cands(self, cands):
        wjson(os.path.join(self.dir, 'cands.json'), cands)


def public_cands(cands):
    keep = ('id', 'w', 'h', 'menu', 'alt', 'kind', 'page', 'src', 'parent')
    return [{k: c.get(k) for k in keep} for c in cands]


from engine.webcollect import prefer_nums, prefer_tokens  # noqa: E402  (웹 버전과 같이 씀)

IMAGE_SLOTS = ('location', 'birdseye', 'siteplan', 'community', 'mh_map')
UNIT_KINDS = ('basic', 'ext', 'iso')


def run_collect(job, draft, home):
    from engine.webcollect import Collector
    job.log(f'분양 홈페이지 읽는 중: {home}')
    c = Collector(home, job.img_dir, draft['types'], log=job.log, prefer=prefer_nums(draft),
                  prefer_tokens=prefer_tokens(draft))
    r = c.run()
    for e in r['errors'][:5]:
        log.info('%s collect: %s', job.id, e)
    old = [x for x in job.cands() if x.get('kind') == 'upload']
    job.save_cands(old + r['cands'])
    wjson(os.path.join(job.dir, 'webtext.json'), {'text': r.get('text', ''), 'pages': r['pages']})
    auto = r['auto']
    imgs = {k: auto.get(k) or [] for k in IMAGE_SLOTS}
    imgs['units'] = {a: {k: v for k, v in u.items() if k in UNIT_KINDS} for a, u in (auto.get('units') or {}).items()}
    job.log(f"그림 후보 {len(r['cands'])}장 (자동 배정: 위치도 {len(imgs['location'])}, 조감도 {len(imgs['birdseye'])}, "
            f"배치도 {len(imgs['siteplan'])}, 커뮤니티 {len(imgs['community'])}, 견본주택 위치도 {len(imgs['mh_map'])}, "
            f"평면도 {len(imgs['units'])}타입)")
    if not imgs['mh_map']:
        job.log('견본주택 위치도(O6)를 홈페이지에서 찾지 못했습니다. 화면 5번에서 넣지 않으면 양식 그림을 그대로 둡니다.')
    if not r['cands']:
        job.log('홈페이지에서 그림을 찾지 못했습니다. 화면에서 그림 파일을 직접 넣어 주세요.')
    return imgs, r.get('text', '')


def check_template(path):
    """양식에 분양개요 시트·공급개요 표가 있는지 가볍게 확인"""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception as e:
        raise UserError(f'양식 엑셀을 열 수 없습니다: {e}')
    names = wb.sheetnames
    ws = wb['분양개요'] if '분양개요' in names else wb[names[0]]
    found = {'공급개요': False, '소계': False, '단위세대': False}
    for row in ws.iter_rows(min_row=1, max_row=300, max_col=6, values_only=True):
        for v in row:
            t = re.sub(r'\s', '', str(v)) if v is not None else ''
            if t.startswith('3.공급개요'):
                found['공급개요'] = True
            if t == '소계':
                found['소계'] = True
            if re.match(r'^\d\.단위세대', t):
                found['단위세대'] = True
    wb.close()
    if not found['공급개요'] or not found['소계']:
        raise UserError('양식에서 "3. 공 급 개 요" 표(구분·소계)를 찾지 못했습니다. 분양개요 및 청약결과 양식(.xlsx)을 넣어 주세요.')
    return {'sheets': names, 'has_result': '청약결과' in names, **found}


def analyze_job(job, notice_name, home_override):
    from engine.notice import NoticeParser, NoticeError
    from engine.options import read_options
    from engine.draft import make_draft, finalize
    try:
        job.set(status='analyzing')
        job.log('모집공고 읽는 중')
        t = time.time()
        np_ = NoticeParser(open(os.path.join(job.dir, 'notice.pdf'), 'rb').read(), notice_name)
        try:
            parsed = np_.run()
        except NoticeError as e:
            raise UserError(str(e))
        if not parsed['types']:
            raise UserError('공고에서 공급대상 표(주택형·세대수)를 찾지 못했습니다.')
        groups = read_options(np_, parsed['types'])
        draft = make_draft(parsed, groups, notice_name)
        draft['complex_name'] = parsed['info'].get('name') or ''
        job.log(f"공고 분석 끝 ({time.time() - t:.0f}초): 타입 {len(draft['types'])}개, 총 {draft['households']:,}세대, "
                f"유상옵션 품목 {len(groups)}개")
        for w in parsed['warnings']:
            job.log('⚠ ' + w)
        # 홈페이지 주소는 화면에서 직접 넣은 것만 쓴다 (2026-10-15 사용자 지정 — 공고 속 주소는 쓰지 않음)
        draft['homepage'] = home_override or ''
        tinfo = rjson(os.path.join(job.dir, 'template.json'), {})
        draft['template'] = tinfo
        wjson(os.path.join(job.dir, 'draft.json'), draft)
        # 홈페이지 그림 (주차대수는 화면에서 직접 입력 — 2026-10-15 사용자 지정)
        imgs = {**{k: [] for k in IMAGE_SLOTS}, 'units': {}}
        if draft.get('homepage'):
            try:
                imgs, _text = run_collect(job, draft, draft['homepage'])
            except Exception as e:
                job.log(f'홈페이지 읽기 실패: {e}')
        else:
            job.log('홈페이지 주소를 넣지 않아 그림은 건너뜁니다. 아래에서 주소를 넣고 "그림 다시 수집"을 누르거나 그림 파일을 직접 넣어 주세요.')
        draft['images'] = imgs
        draft = finalize(draft)
        wjson(os.path.join(job.dir, 'draft.json'), draft)
        job.set(status='ready')
        job.log('분석 끝 — 아래에서 값과 그림을 확인한 뒤 "엑셀 만들기"를 누르세요.')
    except UserError as e:
        job.set(status='error', error=str(e))
        job.log('✖ ' + str(e))
    except Exception as e:
        log.error('analyze %s\n%s', job.id, traceback.format_exc())
        job.set(status='error', error=f'분석 중 오류: {e}')
        job.log(f'✖ 분석 중 오류: {e}')


def build_job(job, draft):
    from engine.draft import finalize
    from engine.build_xlsx import build_workbook
    try:
        job.set(status='building', result=None)
        draft = finalize(draft)
        wjson(os.path.join(job.dir, 'draft.json'), draft)
        cands = {c['id']: c for c in job.cands()}
        def files(ids):
            return [cands[i]['file'] for i in ids or [] if i in cands and os.path.exists(cands[i]['file'])]
        im = draft.get('images') or {}
        images = {k: files(im.get(k)) for k in IMAGE_SLOTS}
        images['units'] = {}
        for a, u in (im.get('units') or {}).items():
            images['units'][a] = {k: (files([v]) or [None])[0] for k, v in u.items() if k in UNIT_KINDS}
        name = safe_filename(draft.get('output_name') or '분양개요 및 청약결과.xlsx')
        out_dir = os.path.join(job.dir, 'out')
        shutil.rmtree(out_dir, ignore_errors=True)
        os.makedirs(out_dir)
        out = os.path.join(out_dir, 'result.xlsx')
        pdf = os.path.join(out_dir, 'preview.pdf')
        if not EXCEL_LOCK.acquire(timeout=1):
            job.log('다른 엑셀 작업이 끝나기를 기다리는 중…')
            EXCEL_LOCK.acquire()
        try:
            job.log('엑셀 작성 중 (이 PC 의 Excel 사용, 1~2분)')
            sec = build_workbook(os.path.join(job.dir, 'template.xlsx'), out, draft, images, log=job.log, preview_pdf=pdf)
        finally:
            EXCEL_LOCK.release()
        pages = []
        try:
            import pymupdf
            doc = pymupdf.open(pdf)
            for i, pg in enumerate(doc):
                fn = os.path.join(out_dir, f'p{i + 1}.png')
                pg.get_pixmap(matrix=pymupdf.Matrix(1.25, 1.25)).save(fn)
                pages.append(i + 1)
        except Exception as e:
            job.log(f'미리보기 만들기 실패: {e}')
        job.set(status='done', result={'name': name, 'pages': pages, 'sec': round(sec)})
        job.log(f'엑셀 완성 ({sec:.0f}초): {name}')
    except Exception as e:
        log.error('build %s\n%s', job.id, traceback.format_exc())
        job.set(status='ready', result=None, error=f'엑셀 작성 중 오류: {e}')
        job.log(f'✖ 엑셀 작성 중 오류: {e}')


def safe_filename(name):
    name = re.sub(r'[\\/:*?"<>|]', '_', name).strip() or '분양개요 및 청약결과.xlsx'
    if not name.lower().endswith('.xlsx'):
        name += '.xlsx'
    return name


# ------------------------------------------------------------------ HTTP

class Handler(BaseHTTPRequestHandler):
    server_version = 'BunyangOverview/1.0'

    def log_message(self, fmt, *args):
        if '/api/status' in (args[0] if args else ''):
            return
        log.info('%s %s', self.client_address[0], fmt % args)

    def is_local(self):
        return self.client_address[0] in ('127.0.0.1', '::1')

    def authorized(self):
        if self.is_local():
            return True
        q = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
        if q.get('k', [''])[0] == KEY or self.headers.get('X-Key') == KEY:
            return True
        return f'k={KEY}' in self.headers.get('Cookie', '')

    def send(self, code, body, ctype='application/json; charset=utf-8', extra=None):
        if isinstance(body, (dict, list)):
            body = json.dumps(body, ensure_ascii=False, default=str).encode('utf-8')
        elif isinstance(body, str):
            body = body.encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', ctype)
        self.send_header('Content-Length', str(len(body)))
        self.send_header('Cache-Control', 'no-store')
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def fail(self, msg, code=400):
        self.send(code, {'ok': False, 'error': msg})

    def body(self):
        n = int(self.headers.get('Content-Length') or 0)
        if n > MAX_UPLOAD:
            raise UserError('파일이 너무 큽니다 (80MB 이하).')
        return self.rfile.read(n) if n else b''

    def q(self):
        return {k: v[0] for k, v in urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query).items()}

    # ---------------------------------------------------------------- GET
    def do_GET(self):
        url = urllib.parse.urlparse(self.path)
        if not self.authorized():
            return self.send(403, '<meta charset="utf-8"><p style="font:16px sans-serif;padding:40px">'
                                  '접속키가 필요합니다. 담당자에게 받은 주소(…?k=…)로 들어와 주세요.</p>',
                             'text/html; charset=utf-8')
        try:
            if url.path in ('/', '/index.html'):
                extra = {}
                if self.q().get('k') == KEY:
                    extra['Set-Cookie'] = f'k={KEY}; Path=/; Max-Age=31536000; SameSite=Lax'
                with open(os.path.join(APP, 'static', 'index.html'), 'rb') as f:
                    return self.send(200, f.read(), 'text/html; charset=utf-8', extra)
            if url.path == '/api/ping':
                return self.send(200, {'ok': True, 'app': 'bunyang-overview', 'local': self.is_local()})
            if url.path == '/api/info':
                info = {'ok': True, 'local': self.is_local()}
                if self.is_local():
                    info['share_url'] = f'http://{lan_ip()}:{PORT}/?k={KEY}'
                return self.send(200, info)
            q = self.q()
            if url.path == '/api/status':
                job = Job(q.get('job'))
                st = job.state()
                out = {'ok': True, **st}
                if q.get('full') == '1' and st.get('status') in ('ready', 'building', 'done'):
                    out['draft'] = rjson(os.path.join(job.dir, 'draft.json'))
                    out['cands'] = public_cands(job.cands())
                return self.send(200, out)
            if url.path == '/api/img':
                job = Job(q.get('job'))
                cid = q.get('id', '')
                if not re.fullmatch(r'[a-z0-9]{6,20}', cid):
                    return self.fail('그림 번호가 올바르지 않습니다.', 404)
                fn = os.path.join(job.img_dir, cid + ('_t.jpg' if q.get('t') == '1' else '.jpg'))
                if not os.path.exists(fn):
                    return self.fail('그림이 없습니다.', 404)
                with open(fn, 'rb') as f:
                    return self.send(200, f.read(), 'image/jpeg', {'Cache-Control': 'max-age=86400'})
            if url.path == '/api/preview':
                job = Job(q.get('job'))
                p = int(q.get('p', '1'))
                fn = os.path.join(job.dir, 'out', f'p{p}.png')
                if not os.path.exists(fn):
                    return self.fail('미리보기가 없습니다.', 404)
                with open(fn, 'rb') as f:
                    return self.send(200, f.read(), 'image/png')
            if url.path == '/api/download':
                job = Job(q.get('job'))
                st = job.state()
                fn = os.path.join(job.dir, 'out', 'result.xlsx')
                if not os.path.exists(fn) or not st.get('result'):
                    return self.fail('아직 만든 파일이 없습니다.', 404)
                name = st['result']['name']
                with open(fn, 'rb') as f:
                    data = f.read()
                return self.send(200, data, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                 {'Content-Disposition': "attachment; filename*=UTF-8''" + urllib.parse.quote(name)})
            return self.fail('없는 주소입니다.', 404)
        except UserError as e:
            return self.fail(str(e), 404)
        except Exception as e:
            log.error('GET %s\n%s', self.path, traceback.format_exc())
            return self.fail(f'오류: {e}', 500)

    # ---------------------------------------------------------------- POST
    def do_POST(self):
        url = urllib.parse.urlparse(self.path)
        if not self.authorized():
            return self.fail('접속키가 필요합니다.', 403)
        try:
            if url.path == '/api/analyze':
                return self.api_analyze()
            if url.path == '/api/build':
                return self.api_build()
            if url.path == '/api/crop':
                return self.api_crop()
            if url.path == '/api/upload_img':
                return self.api_upload_img()
            if url.path == '/api/recollect':
                return self.api_recollect()
            if url.path == '/api/shutdown' and self.is_local():
                self.send(200, {'ok': True})
                threading.Thread(target=self.server.shutdown, daemon=True).start()
                return
            return self.fail('없는 주소입니다.', 404)
        except UserError as e:
            return self.fail(str(e))
        except Exception as e:
            log.error('POST %s\n%s', self.path, traceback.format_exc())
            return self.fail(f'처리 중 오류가 났습니다: {e}', 500)

    def api_analyze(self):
        ctype = self.headers.get('Content-Type', '')
        if 'multipart/form-data' not in ctype:
            raise UserError('파일을 올려 주세요.')
        files, form = parse_multipart(ctype, self.body())
        if 'notice' not in files or 'template' not in files:
            raise UserError('모집공고(PDF)와 분양개요 양식(xlsx)을 모두 올려 주세요.')
        nname, ndata = files['notice']
        tname, tdata = files['template']
        if not nname.lower().endswith('.pdf'):
            raise UserError('모집공고는 PDF 파일이어야 합니다.')
        if not tname.lower().endswith(('.xlsx', '.xlsm')):
            raise UserError('양식은 엑셀(.xlsx) 파일이어야 합니다.')
        cleanup_jobs()
        jid = os.urandom(6).hex()
        d = os.path.join(JOBS, jid)
        os.makedirs(d)
        with open(os.path.join(d, 'notice.pdf'), 'wb') as f:
            f.write(ndata)
        with open(os.path.join(d, 'template.xlsx'), 'wb') as f:
            f.write(tdata)
        try:
            tinfo = check_template(os.path.join(d, 'template.xlsx'))
        except UserError:
            shutil.rmtree(d, ignore_errors=True)
            raise
        tinfo['name'] = tname
        wjson(os.path.join(d, 'template.json'), tinfo)
        job = Job(jid)
        job.set(status='queued', log=[], names={'notice': nname, 'template': tname}, created=time.time())
        home = (form.get('homepage') or '').strip()
        threading.Thread(target=analyze_job, args=(job, nname, home), daemon=True).start()
        log.info('analyze %s %s + %s', jid, nname, tname)
        self.send(200, {'ok': True, 'job': jid})

    def api_build(self):
        req = json.loads(self.body() or b'{}')
        job = Job(req.get('job'))
        if job.state().get('status') in ('building', 'analyzing'):
            raise UserError('지금 처리 중입니다. 끝난 뒤 다시 눌러 주세요.')
        draft = req.get('draft')
        if not isinstance(draft, dict) or not draft.get('types'):
            raise UserError('작성할 값이 없습니다.')
        job.set(status='building', result=None, error=None)
        threading.Thread(target=build_job, args=(job, draft), daemon=True).start()
        self.send(200, {'ok': True})

    def api_crop(self):
        from engine.webcollect import crop_save
        req = json.loads(self.body() or b'{}')
        job = Job(req.get('job'))
        cands = job.cands()
        src = next((c for c in cands if c['id'] == req.get('id')), None)
        if not src:
            raise UserError('그림을 찾지 못했습니다.')
        x0, y0, x1, y1 = [float(v) for v in req.get('box', [0, 0, 1, 1])]
        box = (max(0, round(min(x0, x1) * src['w'])), max(0, round(min(y0, y1) * src['h'])),
               min(src['w'], round(max(x0, x1) * src['w'])), min(src['h'], round(max(y0, y1) * src['h'])))
        if box[2] - box[0] < 20 or box[3] - box[1] < 20:
            raise UserError('자를 영역이 너무 작습니다.')
        c = crop_save(src, box, job.img_dir, cands, 'crop', '자름')
        job.save_cands(cands)
        self.send(200, {'ok': True, 'cand': public_cands([c])[0]})

    def api_upload_img(self):
        from PIL import Image
        import io
        import hashlib
        files, form = parse_multipart(self.headers.get('Content-Type', ''), self.body())
        job = Job(form.get('job'))
        if 'file' not in files:
            raise UserError('그림 파일을 넣어 주세요.')
        name, data = files['file']
        try:
            img = Image.open(io.BytesIO(data))
            img.load()
        except Exception:
            raise UserError('그림 파일(jpg, png)이 아닙니다.')
        from engine.webcollect import Collector
        cid = 'u' + hashlib.md5(data).hexdigest()[:10]
        col = Collector('', job.img_dir)
        path = col.save(img, cid)
        cands = [c for c in job.cands() if c['id'] != cid]
        c = {'id': cid, 'file': path, 'w': img.width, 'h': img.height, 'src': name, 'page': '', 'menu': '직접 넣음',
             'title': '', 'alt': name, 'kind': 'upload'}
        cands.append(c)
        job.save_cands(cands)
        self.send(200, {'ok': True, 'cand': public_cands([c])[0]})

    def api_recollect(self):
        req = json.loads(self.body() or b'{}')
        job = Job(req.get('job'))
        home = (req.get('url') or '').strip()
        if not home:
            raise UserError('홈페이지 주소를 넣어 주세요.')
        draft = rjson(os.path.join(job.dir, 'draft.json'))
        if not draft:
            raise UserError('분석이 끝난 뒤 눌러 주세요.')
        def work():
            try:
                job.set(status='analyzing')
                draft['homepage'] = home
                imgs, text = run_collect(job, draft, home)
                draft['images'] = imgs
                wjson(os.path.join(job.dir, 'draft.json'), draft)
            except Exception as e:
                job.log(f'홈페이지 읽기 실패: {e}')
            job.set(status='ready')
        job.set(status='analyzing')
        threading.Thread(target=work, daemon=True).start()
        self.send(200, {'ok': True})



def main():
    local_only = '--local' in sys.argv
    host = '127.0.0.1' if local_only else '0.0.0.0'
    try:
        srv = ThreadingHTTPServer((host, PORT), Handler)
    except OSError:
        print(f'{PORT} 번 포트를 이미 쓰고 있습니다. 서버가 이미 켜져 있을 수 있습니다.')
        log.warning('port %s busy', PORT)
        sys.exit(2)
    srv.daemon_threads = True
    lines = [f'http://localhost:{PORT}/', '' if local_only else f'http://{lan_ip()}:{PORT}/?k={KEY}']
    with open(URLFILE, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    log.info('start %s', lines)
    try:
        print('분양개요 자동작성 서버가 켜졌습니다.')
        print('  이 PC  :', lines[0])
        if lines[1]:
            print('  다른 PC:', lines[1])
    except Exception:
        pass
    try:
        srv.serve_forever()
    finally:
        try:
            os.remove(URLFILE)
        except OSError:
            pass
        log.info('stop')


if __name__ == '__main__':
    main()
